"""TXT and simple DOCX import/export (not formatting-preserving)."""
from pathlib import Path


def read_text(path: str) -> str:
    file = Path(path)
    if file.suffix.lower() == ".txt":
        return file.read_text(encoding="utf-8-sig")
    if file.suffix.lower() == ".docx":
        from docx import Document
        return "\n".join(p.text for p in Document(str(file)).paragraphs)
    raise ValueError("Поддерживаются только .txt и .docx")


def write_text(path: str, text: str, rtl: bool = False) -> None:
    file = Path(path)
    if file.suffix.lower() == ".txt":
        file.write_text(text, encoding="utf-8-sig")
    elif file.suffix.lower() == ".docx":
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        doc = Document()
        for line in text.split("\n"):
            paragraph = doc.add_paragraph(line)
            if rtl:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                pPr = paragraph._p.get_or_add_pPr()
                bidi = OxmlElement('w:bidi')
                bidi.set(qn('w:val'), '1')
                pPr.append(bidi)
        doc.save(str(file))
    else:
        raise ValueError("Поддерживаются только .txt и .docx")
